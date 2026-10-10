"""Tiny CV files built in code, so the tests don't need binary fixtures."""

import io
import zipfile

from docx import Document

SKILLS = "Python, FastAPI, PostgreSQL, Docker, REST APIs, Git, unit testing"
# Long enough per page that extraction counts as good, not "low text".
LONG_PAGE = "\n".join(
    [
        "Ada Lovelace - Backend Engineer",
        f"Skills: {SKILLS}",
        *[f"Built and shipped feature {i} with Python and PostgreSQL for a team of five." for i in range(8)],
    ]
)


def _escape(line: str) -> str:
    return line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(pages: list[str]) -> bytes:
    """A valid PDF with one text page per item (an empty string = a blank page)."""
    count = len(pages)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(count))
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {count} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for i, text in enumerate(pages):
        shown = " ".join(f"({_escape(line)}) '" for line in text.split("\n")) if text else ""
        stream = f"BT /F1 10 Tf 40 760 Td 14 TL {shown} ET" if text else ""
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>".encode()
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream".encode())

    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table:
        grid = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def make_zip(files: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return buffer.getvalue()
