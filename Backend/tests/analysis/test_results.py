"""Results features: corrections, rewrite suggestions, downloads, comparison."""

import io
from datetime import timedelta

import pytest
from docx import Document
from sqlmodel import select

from app.analysis import export, pipeline
from app.analysis import router as analysis_router
from app.analysis.models import Analysis
from app.analysis.pipeline import ExtractedRequirements, MatchOutput, SuggestOutput
from app.analysis.scoring import apply_rewrites, locate, new_terms, rescore
from app.analysis.service import INTERRUPTED
from app.core.llm import LLMError
from app.core.time import utc_now
from tests.analysis.test_analysis import REQUIREMENTS, VERDICTS, FakeLLM, analyse, new_cv, new_job
from tests.cvs.test_cvs import signed_in

# Lines of tests.cvs.files.LONG_PAGE, the CV every test uploads.
LINE_0 = "Built and shipped feature 0 with Python and PostgreSQL for a team of five."
LINE_1 = "Built and shipped feature 1 with Python and PostgreSQL for a team of five."

SUGGESTIONS = [
    # 2 = FastAPI (partial): reworded, nothing new.
    {"number": 2, "original": LINE_0, "rewritten": LINE_0.replace("Python and", "Python, FastAPI and"), "explanation": "Names FastAPI."},
    # 4 = AWS (matched, but its quote was invented): adds a term the CV lacks.
    {"number": 4, "original": LINE_1, "rewritten": LINE_1.replace(" for", " on AWS for"), "explanation": "Mentions cloud."},
    # 3 = Kubernetes (missing): nothing related, a gap.
    {"number": 3, "original": "", "rewritten": "", "explanation": "Get hands-on with Kubernetes."},
    # Not asked for (1 earned full points), and a repeat: both ignored.
    {"number": 1, "original": LINE_0, "rewritten": "x", "explanation": "?"},
    {"number": 2, "original": "Led a team of 50", "rewritten": "Led a big team", "explanation": "?"},
]


class FakeLLMWithSuggestions(FakeLLM):
    def __init__(self):
        super().__init__()
        self.suggestions = SUGGESTIONS

    def __call__(self, system, user, schema):
        if schema is SuggestOutput:
            self.calls.append(schema.__name__)
            if self.error:
                raise self.error
            return SuggestOutput(suggestions=self.suggestions)
        return super().__call__(system, user, schema)


@pytest.fixture
def fake_llm(monkeypatch):
    fake = FakeLLMWithSuggestions()
    monkeypatch.setattr(pipeline, "chat_json", fake)
    return fake


@pytest.fixture
def auth(client):
    return signed_in(client)


@pytest.fixture
def done(client, auth, fake_llm) -> str:
    """A finished analysis (score 43: see test_analysis.py)."""
    return analyse(client, auth, new_cv(client, auth), new_job(client, auth)).json()["id"]


def correct(client, headers, analysis_id, number, status, note=None):
    return client.patch(f"/analyses/{analysis_id}/requirements/{number}", headers=headers, json={"status": status, "note": note})


def get(client, headers, analysis_id):
    return client.get(f"/analyses/{analysis_id}", headers=headers).json()


# --- corrections ------------------------------------------------------------

def test_correction_rescores_and_keeps_the_model_verdict(client, auth, done):
    # AWS: the model said matched with an invented quote, so 0 points.
    r = correct(client, auth, done, 4, "matched", "I ran our staging on AWS EC2.")
    assert r.status_code == 200
    row = r.json()["result"]["requirements"][3]
    assert (row["status"], row["user_status"], row["user_note"], row["points"]) == ("matched", "matched", "I ran our staging on AWS EC2.", 2)
    assert r.json()["score"] == 71  # 5 of 7
    assert r.json()["result"]["model_score"] == 43
    assert client.get("/analyses", headers=auth).json()[0]["score"] == 71


def test_undoing_a_correction_restores_the_score(client, auth, done):
    correct(client, auth, done, 3, "partial")  # Kubernetes, preferred: +0.5
    assert get(client, auth, done)["score"] == 50
    r = correct(client, auth, done, 3, None, "ignored")
    row = r.json()["result"]["requirements"][2]
    assert (row["user_status"], row["user_note"], r.json()["score"]) == (None, None, 43)


def test_choosing_the_models_own_verdict_is_no_correction(client, auth, done):
    row = correct(client, auth, done, 1, "matched", "yes").json()["result"]["requirements"][0]
    assert (row["user_status"], row["user_note"]) == (None, None)


def test_correction_errors(client, auth, done, monkeypatch):
    assert correct(client, auth, done, 0, "matched").status_code == 404
    assert correct(client, auth, done, 5, "matched").status_code == 404
    assert correct(client, auth, done, 1, "excellent").status_code == 422
    assert correct(client, auth, done, 1, "matched", "x" * 501).status_code == 422
    grace = signed_in(client, email="grace@example.com")
    assert correct(client, grace, done, 1, "missing").status_code == 404

    monkeypatch.setattr(analysis_router, "run_analysis", lambda *args: None)  # stays queued
    queued = analyse(client, auth, new_cv(client, auth, "q.pdf"), new_job(client, auth)).json()["id"]
    assert correct(client, auth, queued, 1, "matched").status_code == 409


# --- suggestions ------------------------------------------------------------

def request(client, headers, analysis_id):
    return client.post(f"/analyses/{analysis_id}/suggestions", headers=headers)


def test_suggestions_are_checked_before_use(client, auth, done, fake_llm):
    r = request(client, auth, done)
    assert r.status_code == 202 and r.json()["suggestions_status"] == "queued"

    body = get(client, auth, done)
    assert body["suggestions_status"] == "done"
    assert body["suggestions_prompt_version"] == pipeline.SUGGESTIONS_PROMPT_VERSION
    by_number = {s["number"]: s for s in body["suggestions"]}
    assert sorted(by_number) == [2, 3, 4]  # only what lost points; repeats dropped

    fastapi, kubernetes, aws = by_number[2], by_number[3], by_number[4]
    assert fastapi["applicable"] and fastapi["new_terms"] == [] and not fastapi["accepted"]
    assert aws["applicable"] and aws["new_terms"] == ["AWS"]  # the user must vouch for it
    assert (kubernetes["original"], kubernetes["applicable"]) == ("", False)  # a gap


def test_invented_original_cannot_be_applied(client, auth, done, fake_llm):
    fake_llm.suggestions = [{"number": 2, "original": "Led a team of 50 engineers", "rewritten": "Led 50 engineers", "explanation": "."}]
    request(client, auth, done)
    suggestion = get(client, auth, done)["suggestions"][0]
    assert suggestion["applicable"] is False
    r = client.patch(f"/analyses/{done}/suggestions/2", headers=auth, json={"accepted": True})
    assert r.status_code == 409


def test_accepting_and_downloading_the_rewritten_cv(client, auth, done, fake_llm):
    request(client, auth, done)
    assert client.get(f"/analyses/{done}/cv-rewrite", headers=auth).status_code == 409  # nothing accepted

    r = client.patch(f"/analyses/{done}/suggestions/2", headers=auth, json={"accepted": True})
    assert r.status_code == 200 and r.json()["suggestions"][0]["accepted"] is True
    assert client.patch(f"/analyses/{done}/suggestions/3", headers=auth, json={"accepted": True}).status_code == 409  # gap
    assert client.patch(f"/analyses/{done}/suggestions/9", headers=auth, json={"accepted": True}).status_code == 404

    r = client.get(f"/analyses/{done}/cv-rewrite", headers=auth)
    assert r.status_code == 200
    assert r.headers["content-type"] == analysis_router.DOCX
    assert 'filename="cv-rewritten.docx"' in r.headers["content-disposition"]
    text = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    assert "with Python, FastAPI and PostgreSQL" in text
    assert LINE_0 not in text  # replaced
    assert LINE_1 in text  # not accepted, unchanged


def test_redoing_suggestions_starts_over(client, auth, done, fake_llm):
    request(client, auth, done)
    client.patch(f"/analyses/{done}/suggestions/2", headers=auth, json={"accepted": True})
    request(client, auth, done)
    assert not any(s["accepted"] for s in get(client, auth, done)["suggestions"])


def test_one_suggestions_run_at_a_time(client, auth, done, monkeypatch):
    monkeypatch.setattr(analysis_router, "run_suggestions", lambda *args: None)  # stays queued
    assert request(client, auth, done).status_code == 202
    assert request(client, auth, done).status_code == 409


def test_interrupted_suggestions_are_marked_failed(client, auth, done, session, monkeypatch):
    monkeypatch.setattr(analysis_router, "run_suggestions", lambda *args: None)
    request(client, auth, done)
    analysis = session.exec(select(Analysis)).one()
    analysis.suggestions_requested_at = utc_now() - timedelta(hours=1)
    session.add(analysis)
    session.commit()
    body = get(client, auth, done)
    assert (body["suggestions_status"], body["suggestions_error"]) == ("failed", INTERRUPTED)
    assert request(client, auth, done).status_code == 202  # can be asked for again


def test_suggestions_failure(client, auth, done, fake_llm):
    fake_llm.error = LLMError("The analysis took too long. Try again later.")
    request(client, auth, done)
    body = get(client, auth, done)
    assert (body["suggestions_status"], body["suggestions_error"]) == ("failed", "The analysis took too long. Try again later.")


def test_nothing_to_improve_skips_the_model(client, auth, fake_llm):
    fake_llm.requirements = REQUIREMENTS[:1]
    fake_llm.verdicts = VERDICTS[:1]  # Python, matched and verified: full points
    analysis_id = analyse(client, auth, new_cv(client, auth), new_job(client, auth)).json()["id"]
    request(client, auth, analysis_id)
    assert get(client, auth, analysis_id)["suggestions"] == []
    assert "SuggestOutput" not in fake_llm.calls


def test_suggestions_need_a_finished_analysis(client, auth, monkeypatch):
    monkeypatch.setattr(analysis_router, "run_analysis", lambda *args: None)
    queued = analyse(client, auth, new_cv(client, auth), new_job(client, auth)).json()["id"]
    assert request(client, auth, queued).status_code == 409


# --- report downloads -------------------------------------------------------

def test_pdf_report(client, auth, done):
    r = client.get(f"/analyses/{done}/report", headers=auth)
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF")
    assert 'filename="ResumeRadar-Backend-at-Acme.pdf"' in r.headers["content-disposition"]


def test_word_report_shows_corrections(client, auth, done):
    correct(client, auth, done, 4, "matched", "I ran staging on AWS.")
    r = client.get(f"/analyses/{done}/report?format=docx", headers=auth)
    assert r.headers["content-type"] == analysis_router.DOCX
    doc = Document(io.BytesIO(r.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
    assert "71% match: Backend at Acme" in text
    assert "the model's own score was 43%" in text
    assert "corrected; model said matched" in cells and "I ran staging on AWS." in cells


def test_report_errors(client, auth, done, monkeypatch):
    assert client.get(f"/analyses/{done}/report?format=txt", headers=auth).status_code == 422
    grace = signed_in(client, email="grace@example.com")
    assert client.get(f"/analyses/{done}/report", headers=grace).status_code == 404
    monkeypatch.setattr(analysis_router, "run_analysis", lambda *args: None)
    queued = analyse(client, auth, new_cv(client, auth, "q.pdf"), new_job(client, auth)).json()["id"]
    assert client.get(f"/analyses/{queued}/report", headers=auth).status_code == 409


def test_pdf_text_survives_unusual_characters():
    assert export._pdf_text(" Python → Go <b>&") == "• Python -&gt; Go &lt;b&gt;&amp;"
    assert export._pdf_text("日本") == "??"


def test_download_names_are_safe():
    assert export.filename("ResumeRadar: Engineer @ Acme/Labs", "pdf") == "ResumeRadar-Engineer-Acme-Labs.pdf"
    assert export.filename("!!!", "docx") == "resumeradar.docx"


# --- compare ----------------------------------------------------------------

def test_compare_two_cvs_for_the_same_job(client, auth, fake_llm):
    job_id = new_job(client, auth)
    first = analyse(client, auth, new_cv(client, auth, "a.pdf"), job_id).json()["id"]
    second = analyse(client, auth, new_cv(client, auth, "b.pdf"), job_id).json()["id"]
    correct(client, auth, second, 4, "matched")

    r = client.get(f"/analyses/compare?first={first}&second={second}", headers=auth)
    assert r.status_code == 200
    body = r.json()
    assert (body["first"]["cv_name"], body["second"]["cv_name"]) == ("a", "b")
    assert (body["first"]["score"], body["second"]["score"]) == (43, 71)
    aws = next(row for row in body["rows"] if row["requirement"] == "AWS")
    assert (aws["first"]["points"], aws["second"]["points"], aws["second"]["corrected"]) == (0, 2, True)
    assert len(body["rows"]) == 4


def test_compare_errors(client, auth, fake_llm):
    first = analyse(client, auth, new_cv(client, auth), new_job(client, auth)).json()["id"]
    other_job = analyse(client, auth, new_cv(client, auth, "b.pdf"), new_job(client, auth)).json()["id"]
    r = client.get(f"/analyses/compare?first={first}&second={other_job}", headers=auth)
    assert r.status_code == 422 and "same job" in r.json()["detail"]
    grace = signed_in(client, email="grace@example.com")
    assert client.get(f"/analyses/compare?first={first}&second={first}", headers=grace).status_code == 404
    assert client.get(f"/analyses/compare?first={first}", headers=auth).status_code == 422


# --- the scoring helpers ----------------------------------------------------

def test_rescore_prefers_corrections_over_evidence():
    rows = [
        {"importance": "required", "status": "matched", "evidence_verified": False, "user_status": None},
        {"importance": "preferred", "status": "missing", "evidence_verified": None, "user_status": "matched"},
    ]
    score, scoring = rescore(rows)
    assert [r["points"] for r in rows] == [0, 1]
    assert (scoring["earned"], scoring["possible"], score) == (1, 3, 33)


@pytest.mark.parametrize(
    ("original", "found"),
    [
        ("Skills: Python, FastAPI", True),
        ("skills python fastapi", True),  # punctuation and case don't matter
        ("Python FastAPI Docker", False),  # but no words may be skipped
        ("", False),
    ],
)
def test_locate(original, found):
    assert (locate(original, "Skills: Python,\nFastAPI, PostgreSQL, Docker") is not None) is found


def test_new_terms_flags_names_numbers_and_job_words():
    cv = "Built REST APIs with FastAPI and Docker."
    rewrite = "Built production REST APIs with FastAPI, Docker and Kubernetes on AWS for 5 years, scaling microservices"
    assert new_terms(rewrite, "Built REST APIs with FastAPI.", cv, "Microservices on Kubernetes") == ["Kubernetes", "AWS", "5", "microservices"]


def test_new_terms_ignores_a_capital_that_starts_the_sentence():
    # From the real run: "Strong" isn't a name; "proficiency" is a claim lifted from the job.
    rewrite = "Strong proficiency with Python and FastAPI for building REST APIs"
    cv = "Python • FastAPI • REST APIs"
    assert new_terms(rewrite, cv, cv, "Strong proficiency with React and Python") == ["proficiency"]
    # A name starting a sentence is still caught when the job asked for it.
    assert new_terms("Deployed on AWS. Kubernetes later", cv, cv, "Kubernetes and cloud") == ["AWS", "Kubernetes"]


def test_apply_rewrites_handles_overlaps_and_order():
    cv = "Line one here.\nLine two here.\nLine three here."
    suggestions = [
        {"accepted": True, "applicable": True, "original": "Line two here.", "rewritten": "Second line."},
        {"accepted": True, "applicable": True, "original": "Line one here.", "rewritten": "First line."},
        {"accepted": True, "applicable": True, "original": "two here. Line three", "rewritten": "overlap"},  # overlaps: skipped
        {"accepted": False, "applicable": True, "original": "Line three here.", "rewritten": "not accepted"},
    ]
    assert apply_rewrites(cv, suggestions) == "First line.\nSecond line.\nLine three here."
