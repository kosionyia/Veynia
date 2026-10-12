"""Downloadable documents: an analysis report (PDF or Word) and the rewritten CV (Word).

The report is described once as a list of blocks, then drawn by each format,
so the PDF and Word versions always say the same thing.
"""

import io
from typing import Literal
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.analysis.schemas import AnalysisDetail
from app.analysis.scoring import effective_status

VERDICT = {"matched": "Matched", "partial": "Partly", "missing": "Missing"}

Block = tuple[Literal["title", "meta", "heading", "para", "table"], object]


def job_label(analysis: AnalysisDetail) -> str:
    if analysis.job_title and analysis.job_company:
        return f"{analysis.job_title} at {analysis.job_company}"
    return analysis.job_title or analysis.job_company or "Untitled job"


def _blocks(analysis: AnalysisDetail) -> list[Block]:
    result = analysis.result
    assert result is not None  # only finished analyses are exported
    finished = analysis.finished_at or analysis.created_at
    blocks: list[Block] = [
        ("title", f"{analysis.score}% match: {job_label(analysis)}"),
        ("meta", f"CV: {analysis.cv_name} · analysed {finished:%d %b %Y} with {analysis.model} (prompt {analysis.prompt_version})"),
        ("para", result.summary),
    ]
    blocks += [("para", f"Note: {warning}") for warning in result.warnings]

    s = result.scoring
    corrected = sum(1 for r in result.requirements if r.user_status)
    how = (
        f"Required requirements are worth {s.weights['required']:g} points, preferred ones "
        f"{s.weights['preferred']:g}. Matched earns all of them, partly half, missing none. A quote "
        f"that can't be found in the CV earns nothing. {s.earned:g} of {s.possible:g} points = {analysis.score}%."
    )
    if corrected:
        how += f" {corrected} verdict(s) were corrected by the user; the model's own score was {result.model_score}%."
    blocks += [("heading", "How the score is calculated"), ("para", how)]

    rows = [["Verdict", "Requirement", "Evidence from the CV", "Points"]]
    for r in result.requirements:
        verdict = VERDICT[effective_status(r.model_dump())]
        if r.user_status:
            verdict += f" (corrected; model said {VERDICT[r.status].lower()})"
        evidence = f"“{r.cv_evidence}”" if r.cv_evidence else ""
        if r.cv_evidence and r.evidence_verified is False:
            evidence += " [not found in the CV]"
        why = r.user_note if r.user_status and r.user_note else r.reason
        rows.append([verdict, f"{r.requirement} ({r.importance})\n{why}", evidence, f"{r.points:g}/{r.max_points:g}"])
    blocks += [("heading", "Requirement by requirement"), ("table", rows)]

    accepted = [s for s in analysis.suggestions or [] if s.accepted]
    if accepted:
        blocks.append(("heading", "Accepted CV rewrites"))
        for s in accepted:
            blocks.append(("para", f"{s.requirement}: “{s.original}” → “{s.rewritten}”"))
    return blocks


# --- PDF --------------------------------------------------------------------

def _pdf_text(text: str) -> str:
    """The built-in PDF fonts only cover Western European characters; swap the
    rest (e.g. PDF bullet glyphs from the CV) rather than print empty boxes."""
    text = text.replace("", "•").replace("→", "->")
    return escape(text.encode("cp1252", errors="replace").decode("cp1252")).replace("\n", "<br/>")


def report_pdf(analysis: AnalysisDetail) -> bytes:
    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=8, leading=10)
    flow = []
    for kind, value in _blocks(analysis):
        if kind == "title":
            flow.append(Paragraph(_pdf_text(str(value)), styles["Title"]))
        elif kind == "meta":
            flow.append(Paragraph(_pdf_text(str(value)), styles["Italic"]))
        elif kind == "heading":
            flow += [Spacer(1, 4 * mm), Paragraph(_pdf_text(str(value)), styles["Heading2"])]
        elif kind == "para":
            flow.append(Paragraph(_pdf_text(str(value)), styles["BodyText"]))
        else:
            rows = [[Paragraph(_pdf_text(cell), small) for cell in row] for row in value]  # type: ignore[union-attr]
            table = Table(rows, colWidths=[30 * mm, 65 * mm, 65 * mm, 18 * mm], repeatRows=1)
            table.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#cccccc")),
            ]))
            flow.append(table)
    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4, title=job_label(analysis), leftMargin=15 * mm, rightMargin=15 * mm).build(flow)
    return buffer.getvalue()


# --- Word -------------------------------------------------------------------

def report_docx(analysis: AnalysisDetail) -> bytes:
    doc = Document()
    for kind, value in _blocks(analysis):
        if kind == "title":
            doc.add_heading(str(value), level=1)
        elif kind == "meta":
            doc.add_paragraph().add_run(str(value)).italic = True
        elif kind == "heading":
            doc.add_heading(str(value), level=2)
        elif kind == "para":
            doc.add_paragraph(str(value))
        else:
            rows = value  # type: ignore[assignment]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))  # type: ignore[arg-type]
            table.style = "Table Grid"
            for r, row in enumerate(rows):  # type: ignore[arg-type]
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell
                    if r == 0:
                        for run in table.cell(r, c).paragraphs[0].runs:
                            run.bold = True
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def cv_docx(text: str) -> bytes:
    """Plain Word file of CV text: one paragraph per line. The original layout
    isn't stored (only the text is), so this is for pasting or restyling."""
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    previous_blank = False
    for line in text.replace("", "•").split("\n"):
        blank = not line.strip()
        if not (blank and previous_blank):
            doc.add_paragraph(line.rstrip())
        previous_blank = blank
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def filename(stem: str, extension: str) -> str:
    """A safe download name: letters, digits and dashes only."""
    cleaned = "".join(c if c.isalnum() else "-" for c in stem).strip("-")
    while "--" in cleaned:
        cleaned = cleaned.replace("--", "-")
    return f"{(cleaned or 'resumeradar')[:80]}.{extension}"

